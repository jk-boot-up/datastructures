from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

BG       = RGBColor(0x1B, 0x21, 0x30)
PANEL    = RGBColor(0x24, 0x2C, 0x40)
PANEL2   = RGBColor(0x2C, 0x35, 0x49)
BORDER   = RGBColor(0x4A, 0x55, 0x70)
TEXT     = RGBColor(0xEE, 0xF1, 0xF7)
MUTED    = RGBColor(0x9A, 0xA6, 0xC2)
ACCENT   = RGBColor(0x4D, 0xA3, 0xFF)
ACCENT2  = RGBColor(0xFF, 0x9D, 0x4D)
GOOD     = RGBColor(0x4F, 0xD9, 0xA0)
CURSOR_C = RGBColor(0xFF, 0xCF, 0x4D)
BAD      = RGBColor(0xFF, 0x6B, 0x81)
ON_ACCENT = RGBColor(0x10, 0x14, 0x1F)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide():
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def add_text(slide, x, y, w, h, text, size=18, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, font="Calibri", italic=False, anchor=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.55),
              title, size=30, color=TEXT, bold=True)
    if subtitle:
        add_text(slide, Inches(0.6), Inches(0.98), Inches(12.1), Inches(0.45),
                  subtitle, size=15, color=MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.55),
                                    Inches(12.1), Pt(2.5))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, gap_pt=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_pt)
        text, kwargs = item if isinstance(item, tuple) else (item, {})
        run = p.add_run()
        run.text = "•  " + text
        run.font.size = Pt(kwargs.get("size", size))
        run.font.color.rgb = kwargs.get("color", color)
        run.font.bold = kwargs.get("bold", False)
        run.font.name = "Calibri"
    return box


def rounded_box(slide, x, y, w, h, fill=PANEL2, line_color=BORDER, line_w=1.75,
                 text=None, text_color=TEXT, size=20, bold=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.18
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_color
    shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if text is not None:
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = text_color
        run.font.name = "Consolas"
    return shp


def pill(slide, x, y, w, h, text, fill, text_color=ON_ACCENT, size=11):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.5
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = text_color
    run.font.name = "Calibri"
    return shp


def connector(slide, x1, y1, x2, y2, color=MUTED, width=2.25):
    conn = slide.shapes.add_connector(1, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.shadow.inherit = False
    return conn


NODE_W, NODE_H = Inches(1.2), Inches(0.8)


def centered_x_start(n, gap):
    total = (n - 1) * gap + NODE_W
    return Emu(int((Inches(13.333) - total) / 2))


def node_row(slide, values, top, x_start=None, gap=Inches(1.9),
             highlight_idx=None, new_idx=None, next_labels=None, tags=None):
    if x_start is None:
        x_start = centered_x_start(len(values), gap)
    for i, v in enumerate(values):
        x = x_start + i * gap
        line_color = BAD if (highlight_idx is not None and i == highlight_idx) \
            else (GOOD if (new_idx is not None and i == new_idx) else BORDER)
        line_w = 3.25 if (i in (highlight_idx, new_idx)) else 1.75
        rounded_box(slide, x, top, NODE_W, NODE_H, fill=PANEL,
                    line_color=line_color, line_w=line_w, text=str(v),
                    text_color=TEXT, size=22)
        add_text(slide, x, top - Inches(0.32), NODE_W, Inches(0.28),
                  f"idx {i}", size=10, color=MUTED, align=PP_ALIGN.CENTER)
        if next_labels is not None:
            nv = next_labels[i]
            add_text(slide, x - Inches(0.15), top + NODE_H + Inches(0.06),
                      NODE_W + Inches(0.3), Inches(0.3),
                      f"next: {nv}", size=12, color=MUTED, align=PP_ALIGN.CENTER)
        if tags and i in tags:
            for j, (label, color) in enumerate(tags[i]):
                py = top - Inches(0.68) - j * Inches(0.42)
                pill(slide, x - Inches(0.05), py, NODE_W + Inches(0.1), Inches(0.32),
                     label, color)


def code_box(slide, x, y, w, h, lines, highlight_idx=None, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.03
    shp.fill.solid(); shp.fill.fore_color.rgb = PANEL2
    shp.line.color.rgb = BORDER
    shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = Emu(137160)
    tf.margin_top = Emu(91440)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = "Consolas"
        run.font.bold = bool(highlight_idx and i in highlight_idx)
        run.font.color.rgb = TEXT if (highlight_idx is None or i not in highlight_idx) else BAD
    return shp


def footer(slide, text):
    add_text(slide, Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.35),
              text, size=11, color=MUTED)


# Slide 1 — Title
s = new_slide()
add_text(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.2),
          "Singly Linked List — reverse()", size=44, bold=True, color=TEXT)
add_text(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.6),
          "An interactive walkthrough of building and reversing a singly linked list",
          size=20, color=MUTED)
add_text(s, Inches(0.9), Inches(4.15), Inches(11.5), Inches(0.5),
          "com.jk.explore.dsa.linkedlists.single.reverse.LinkedList<T>",
          size=15, color=ACCENT, font="Consolas")
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.35), Inches(2.4), Pt(4.5))
line.fill.solid(); line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

# Slide 2 — Agenda
s = new_slide()
add_title(s, "Agenda")
bullets(s, Inches(0.8), Inches(1.9), Inches(11.5), Inches(4.8), [
    "The LinkedList<T> / Node<T> structure",
    "Building a list — addFirst, addLast, add(index, data)",
    "reverse() — control flow and complexity",
    "Frame-by-frame trace of reverse() on [1, 2, 3]",
    "Other operations at a glance",
    "Live demo — linked-list-reverse-demo.html",
], size=20, gap_pt=16)
footer(s, "docs/linked-list-reverse-demo.html   ·   docs/LinkedList-Reverse-Diagrams.md")

# Slide 3 — Structure  (caption moved ABOVE the class boxes; node row pushed
# well clear of the class boxes so idx labels / head-tail tags never collide
# with the caption or the boxes — this was the confirmed overlap bug)
s = new_slide()
add_title(s, "The Structure", "Two classes: a thin wrapper (head/tail/size) around a chain of nodes")
add_text(s, Inches(0.6), Inches(1.68), Inches(8.0), Inches(0.3),
          "A list of [1, 2, 3] before reversal", size=13, color=MUTED)

box1 = rounded_box(s, Inches(1.2), Inches(2.15), Inches(4.6), Inches(2.15), fill=PANEL,
            line_color=ACCENT, line_w=2.25, text="")
tf = box1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run(); run.text = "LinkedList<T>"
run.font.bold = True; run.font.size = Pt(19); run.font.color.rgb = ACCENT; run.font.name = "Consolas"
for line in ["", "- head: Node<T>", "- tail: Node<T>", "- size: int", "",
             "+ addFirst / addLast", "+ add(index, data)", "+ reverse()", "+ isEmpty / getSize"]:
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = line
    r2.font.size = Pt(13); r2.font.color.rgb = TEXT; r2.font.name = "Consolas"

box2 = rounded_box(s, Inches(7.5), Inches(2.5), Inches(3.6), Inches(1.45), fill=PANEL,
            line_color=GOOD, line_w=2.25, text="")
tf2 = box2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
run = p.add_run(); run.text = "Node<T>"
run.font.bold = True; run.font.size = Pt(19); run.font.color.rgb = GOOD; run.font.name = "Consolas"
for line in ["", "- data: T", "- next: Node<T>"]:
    p2 = tf2.add_paragraph()
    r2 = p2.add_run(); r2.text = line
    r2.font.size = Pt(13); r2.font.color.rgb = TEXT; r2.font.name = "Consolas"

connector(s, Inches(5.8), Inches(3.22), Inches(7.5), Inches(3.22), color=MUTED, width=2.25)
add_text(s, Inches(5.85), Inches(2.87), Inches(1.6), Inches(0.3), "links →", size=11, color=MUTED)

node_row(s, ["1", "2", "3"], top=Inches(5.7), gap=Inches(1.7),
         next_labels=["2", "3", "null"],
         tags={0: [("head", ACCENT)], 2: [("tail", ACCENT2)]})
footer(s, "Section 1 — Structure  ·  LinkedList-Reverse-Diagrams.md")

# Slide 4 — reverse() control flow
s = new_slide()
add_title(s, "reverse() — Control Flow", "Classic iterative 3-pointer reversal: prev, cursor, and a saved post")
bullets(s, Inches(0.8), Inches(1.95), Inches(7.1), Inches(4.6), [
    "isEmpty()?  → return (nothing to do)",
    "size == 1?  → return (single node is its own reverse)",
    "Init: cursor = head, prev = null; save oldHead / oldTail",
    "While cursor != null:",
    ("   post = cursor.next", {"size": 15, "color": MUTED}),
    ("   cursor.next = prev", {"size": 15, "color": MUTED}),
    ("   prev = cursor; cursor = post", {"size": 15, "color": MUTED}),
    "head = oldTail; tail = oldHead",
], size=18, gap_pt=12)
box3 = rounded_box(s, Inches(8.3), Inches(2.0), Inches(4.3), Inches(3.4), fill=PANEL, line_color=BORDER,
            line_w=1.75, text="")
tf3 = box3.text_frame
tf3.word_wrap = True
p = tf3.paragraphs[0]
r = p.add_run(); r.text = "Time / Space"
r.font.bold = True; r.font.size = Pt(18); r.font.color.rgb = TEXT
for label, val, why in [("Time", "O(n)", "each node visited once"),
                         ("Space", "O(1)", "fixed pointers, no extra structure")]:
    p2 = tf3.add_paragraph()
    r2 = p2.add_run(); r2.text = f"{label}: "
    r2.font.size = Pt(16); r2.font.color.rgb = MUTED
    r3 = p2.add_run(); r3.text = val
    r3.font.size = Pt(16); r3.font.bold = True; r3.font.color.rgb = GOOD
    p3 = tf3.add_paragraph()
    r4 = p3.add_run(); r4.text = why
    r4.font.size = Pt(13); r4.font.italic = True; r4.font.color.rgb = MUTED
footer(s, "Section 2 — Control flow  ·  LinkedList-Reverse-Diagrams.md")

# Slide 5 — source code
s = new_slide()
add_title(s, "reverse() — Source", "The exact method the demo animates, line by line")
code_lines = [
    "public void reverse() {",
    "    if (isEmpty()) { return; }",
    "    if (size == 1) { return; }",
    "    Node<T> cursorNode = head;",
    "    Node<T> preCursorNode = null;",
    "    Node<T> postCursorNode = null;",
    "    Node<T> oldHeadNode = head;",
    "    Node<T> oldTailNode = tail;",
    "    while (cursorNode != null) {",
    "        postCursorNode = cursorNode.next;",
    "        cursorNode.next = preCursorNode;",
    "        preCursorNode = cursorNode;",
    "        cursorNode = postCursorNode;",
    "    }",
    "    head = oldTailNode;",
    "    tail = oldHeadNode;",
    "}",
]
code_box(s, Inches(1.4), Inches(1.9), Inches(10.5), Inches(4.9), code_lines,
         highlight_idx={9, 10, 11, 12}, size=16)
footer(s, "src/main/java/.../reverse/LinkedList.java")

# Slides 6-10 — Frame-by-frame trace
frames = [
    dict(title="Frame 0 — before the loop starts",
         caption="prev = null, cursor = 1",
         next_labels=["2", "3", "null"], highlight=None,
         tags={0: [("head", ACCENT), ("cursor", CURSOR_C)], 2: [("tail", ACCENT2)]}),
    dict(title="Frame 1 — after processing node 1",
         caption="cursor.next = prev rewired node 1.  prev = 1, cursor = 2",
         next_labels=["null", "3", "null"], highlight=0,
         tags={0: [("head", ACCENT), ("prev", GOOD)], 1: [("cursor", CURSOR_C)], 2: [("tail", ACCENT2)]}),
    dict(title="Frame 2 — after processing node 2",
         caption="prev = 2, cursor = 3",
         next_labels=["null", "1", "null"], highlight=1,
         tags={0: [("head", ACCENT)], 1: [("prev", GOOD)], 2: [("tail", ACCENT2), ("cursor", CURSOR_C)]}),
    dict(title="Frame 3 — after processing node 3 (loop exits)",
         caption="prev = 3, cursor = null — every node has been relinked",
         next_labels=["null", "1", "2"], highlight=2,
         tags={0: [("head", ACCENT)], 2: [("tail", ACCENT2), ("prev", GOOD)]}),
    dict(title="Frame 4 — after head = oldTail; tail = oldHead;",
         caption="head = 3, tail = 1 — next values unchanged, only head/tail swapped",
         next_labels=["null", "1", "2"], highlight=None,
         tags={0: [("tail", ACCENT2)], 2: [("head", ACCENT)]}),
]
for f in frames:
    s = new_slide()
    add_title(s, f["title"])
    add_text(s, Inches(0.8), Inches(1.75), Inches(11.5), Inches(0.4),
              f["caption"], size=16, color=MUTED)
    node_row(s, ["1", "2", "3"], top=Inches(3.5), gap=Inches(2.0),
             highlight_idx=f["highlight"], next_labels=f["next_labels"], tags=f["tags"])
    if f["highlight"] is not None:
        add_text(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
                  "Red border = the next pointer that just changed this frame", size=13,
                  color=BAD, italic=True)
    footer(s, "Section 3 — Frame-by-frame trace  ·  LinkedList-Reverse-Diagrams.md")

# Slide 11 — Result (same physical columns as Frame 4 — nodes never move,
# only head/tail badges did; reversed order is read off head → tail → null)
s = new_slide()
add_title(s, "Result", "Same three nodes, same columns as Frame 4 — only head/tail moved, next pointers now run right-to-left")
node_row(s, ["1", "2", "3"], top=Inches(3.5), gap=Inches(2.0),
         next_labels=["null", "1", "2"],
         tags={0: [("tail", ACCENT2)], 2: [("head", ACCENT)]})
add_text(s, Inches(0.8), Inches(5.9), Inches(11.5), Inches(0.5),
          "Reading from head (idx 2) forward: 3 → 2 → 1 → null.  [1, 2, 3] has logically become [3, 2, 1].",
          size=15, color=TEXT, align=PP_ALIGN.CENTER)
footer(s, "Section 3 — Frame-by-frame trace  ·  LinkedList-Reverse-Diagrams.md")

# Slides 12-14 — Other operations
ops = [
    dict(title="addFirst(0)  —  O(1)", note="Only head moves",
         values=["0", "1", "2", "3"], new_idx=0,
         next_labels=["1", "2", "3", "null"],
         tags={0: [("head", ACCENT), ("new", GOOD)], 3: [("tail", ACCENT2)]}),
    dict(title="addLast(4)  —  O(1)", note="Only tail moves",
         values=["1", "2", "3", "4"], new_idx=3,
         next_labels=["2", "3", "4", "null"],
         tags={0: [("head", ACCENT)], 3: [("tail", ACCENT2), ("new", GOOD)]}),
    dict(title="add(1, 99)  —  O(n)", note="Walks from head to index 1 first; head/tail unchanged",
         values=["1", "99", "2", "3"], new_idx=1,
         next_labels=["99", "2", "3", "null"],
         tags={0: [("head", ACCENT)], 1: [("new", GOOD)], 3: [("tail", ACCENT2)]}),
]
for op in ops:
    s = new_slide()
    add_title(s, op["title"], op["note"])
    node_row(s, op["values"], top=Inches(3.5), gap=Inches(2.0),
             new_idx=op["new_idx"], next_labels=op["next_labels"], tags=op["tags"])
    footer(s, "Section 4 — Other operations at a glance  ·  LinkedList-Reverse-Diagrams.md")

# Slide 15 — Complexity summary
s = new_slide()
add_title(s, "Complexity Summary")
rows = [
    ("Method", "Time", "Space", "Notes"),
    ("LinkedList(T data)", "O(1)", "O(1)", "one node allocated"),
    ("add / addLast", "O(1)", "O(1)", "tail pointer avoids traversal"),
    ("addFirst", "O(1)", "O(1)", "only head moves"),
    ("add(index, data)", "O(n)", "O(1)", "O(1) at the ends, else walks to index"),
    ("reverse()", "O(n)", "O(1)", "single pass, in-place pointer flip"),
    ("isEmpty()", "O(1)", "O(1)", "cached size field"),
    ("toString()", "O(n)", "O(n)", "string proportional to node count"),
]
table_shape = s.shapes.add_table(len(rows), 4, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.6))
table = table_shape.table
widths = [Inches(3.2), Inches(1.4), Inches(1.4), Inches(5.7)]
for i, w in enumerate(widths):
    table.columns[i].width = w
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PANEL2 if r == 0 else PANEL
        cell.margin_left = Emu(91440); cell.margin_right = Emu(91440)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run(); run.text = val
        run.font.size = Pt(13 if r else 14)
        run.font.bold = (r == 0)
        run.font.color.rgb = ACCENT if r == 0 else TEXT
        run.font.name = "Calibri" if c == 3 else "Consolas"
footer(s, "Section 4 — Other operations at a glance  ·  LinkedList-Reverse-Diagrams.md")

# Slide 16 — Live demo script
s = new_slide()
add_title(s, "Live Demo", "linked-list-reverse-demo.html — keyboard: ←/→ to step, space to play/pause")
bullets(s, Inches(0.8), Inches(1.95), Inches(11.7), Inches(4.8), [
    "Open linked-list-reverse-demo.html — starts on Build & Modify with the list [1]",
    "Click + addLast() twice, then + addFirst() once — watch the code panel highlight the branch that ran",
    "Click + add(middle, data) once — complexity flips to O(n), traversal loop highlighted",
    "Switch to Reverse Animation, click \"demo: 5 nodes\"",
    "Press → a few times — call out prev, cursor, and the red \"link being rewritten\" arc",
    "Click ▶ play to run end-to-end, or ⏭ to jump to the final swapped head/tail",
    "For the why: LinkedList-Reverse-Diagrams.md section 3 — the frame table matches exactly",
], size=17, gap_pt=14)
footer(s, "docs/linked-list-reverse-demo.html  ·  docs/LinkedList-Reverse-Diagrams.md")

# Slide 17 — Thank you
s = new_slide()
add_text(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.0),
          "Questions?", size=40, bold=True, color=TEXT)
add_text(s, Inches(0.9), Inches(3.9), Inches(11.5), Inches(0.5),
          "docs/linked-list-reverse-demo.html   ·   docs/LinkedList-Reverse-Diagrams.md",
          size=16, color=ACCENT, font="Consolas")

out_path = "/Users/jk/gitbase/datastructures/dsa-java-maven/linkedlists/singly-linked-list/linkedlist/docs/LinkedList-Reverse-Presentation.pptx"
prs.save(out_path)
print("Saved:", out_path, "slides:", len(prs.slides._sldIdLst))
